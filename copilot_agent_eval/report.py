"""Report generation: Excel, self-contained HTML, and JUnit XML."""
from __future__ import annotations

import html
import json
import logging
import xml.etree.ElementTree as ET
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import TestResult

log = logging.getLogger(__name__)

_STATUS_FILL = {
    "PASS": PatternFill("solid", fgColor="C6EFCE"),
    "FAIL": PatternFill("solid", fgColor="FFC7CE"),
    "TIMEOUT": PatternFill("solid", fgColor="FFEB9C"),
    "ERROR": PatternFill("solid", fgColor="FFC7CE"),
}


def generate_reports(
    results: list[TestResult],
    output_dir: str | Path,
    formats: list[str],
    raw_activities: dict[str, list[dict]] | None = None,
) -> dict[str, Path]:
    out = Path(output_dir)
    out.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    produced: dict[str, Path] = {}

    if "excel" in formats:
        path = out / f"copilot_agent_results_{stamp}.xlsx"
        _write_excel(results, path)
        produced["excel"] = path

    if "html" in formats:
        path = out / f"copilot_agent_report_{stamp}.html"
        _write_html(results, path, raw_activities or {})
        produced["html"] = path

    if "junit" in formats:
        path = out / f"copilot_agent_results_{stamp}.xml"
        _write_junit(results, path)
        produced["junit"] = path

    return produced


# --------------------------------------------------------------------- Excel
def _write_excel(results: list[TestResult], path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Results"

    headers = [
        "Case ID", "Status", "Priority", "Query", "Response",
        "Keyword Hits", "Keyword Misses", "Error Phrases",
        "First Token (ms)", "Total (ms)", "Conversation ID",
        "Timestamp", "Error", "Notes",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for r in results:
        ws.append([
            r.case_id, r.status, "", r.query, r.response_text,
            ", ".join(r.keyword_hits), ", ".join(r.keyword_misses),
            ", ".join(r.error_phrases_hit),
            r.first_token_latency_ms, r.total_latency_ms,
            r.conversation_id, r.timestamp, r.error, "",
        ])
        row_idx = ws.max_row
        fill = _STATUS_FILL.get(r.status)
        if fill:
            ws.cell(row=row_idx, column=2).fill = fill

    widths = [12, 10, 8, 40, 60, 25, 25, 20, 16, 12, 36, 20, 30, 20]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"

    # Summary sheet
    summary = wb.create_sheet("Summary")
    total = len(results)
    passed = sum(1 for r in results if r.status == "PASS")
    failed = sum(1 for r in results if r.status == "FAIL")
    timeout = sum(1 for r in results if r.status == "TIMEOUT")
    error = sum(1 for r in results if r.status == "ERROR")
    summary.append(["Metric", "Value"])
    summary.append(["Total", total])
    summary.append(["Passed", passed])
    summary.append(["Failed", failed])
    summary.append(["Timeout", timeout])
    summary.append(["Error", error])
    summary.append(["Pass rate", f"{(passed / total * 100):.1f}%" if total else "n/a"])
    summary.append(["Generated", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
    for cell in summary[1]:
        cell.font = Font(bold=True)
    summary.column_dimensions["A"].width = 14
    summary.column_dimensions["B"].width = 24

    # Steps sheet — full telemetry timeline per case.
    try:
        from .telemetry import STEP_HEADERS
    except Exception:
        STEP_HEADERS = ["Offset (ms)", "Type", "Name", "Category",
                        "Summary", "Detail", "Activity ID"]
    steps_ws = wb.create_sheet("Steps")
    steps_ws.append(["Case ID"] + STEP_HEADERS)
    for cell in steps_ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for r in results:
        for s in (r.steps or []):
            row = [r.case_id] + (s.as_row() if hasattr(s, "as_row") else [s])
            steps_ws.append(row)
    step_widths = [10, 12, 14, 22, 18, 50, 60, 20]
    for i, w in enumerate(step_widths, start=1):
        steps_ws.column_dimensions[get_column_letter(i)].width = w
    for row in steps_ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    steps_ws.freeze_panes = "A2"

    wb.save(path)
    log.info("Excel report written: %s", path)


# ---------------------------------------------------------------------- HTML
_HTML_TEMPLATE = """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8"/>
<title>Copilot Agent Eval Report</title>
<style>
  body {{ font-family: -apple-system, Segoe UI, Roboto, Helvetica, Arial, sans-serif;
         margin: 24px; color: #1f2937; background: #f9fafb; }}
  h1 {{ margin-bottom: 4px; }}
  .meta {{ color: #6b7280; margin-bottom: 20px; }}
  .summary {{ display: flex; gap: 12px; margin-bottom: 24px; flex-wrap: wrap; }}
  .card {{ background: #fff; border: 1px solid #e5e7eb; border-radius: 8px;
           padding: 12px 18px; min-width: 110px; box-shadow: 0 1px 2px rgba(0,0,0,.04); }}
  .card .n {{ font-size: 26px; font-weight: 700; }}
  .card .l {{ font-size: 12px; color: #6b7280; text-transform: uppercase; letter-spacing: .04em; }}
  .PASS .n {{ color: #059669; }} .FAIL .n, .ERROR .n {{ color: #dc2626; }}
  .TIMEOUT .n {{ color: #d97706; }}
  table {{ width: 100%; border-collapse: collapse; background: #fff;
           border: 1px solid #e5e7eb; border-radius: 8px; overflow: hidden; }}
  th, td {{ padding: 10px 12px; border-bottom: 1px solid #f3f4f6; text-align: left;
            vertical-align: top; font-size: 13px; }}
  th {{ background: #f3f4f6; font-size: 12px; text-transform: uppercase;
        letter-spacing: .04em; color: #374151; }}
  .badge {{ display: inline-block; padding: 2px 8px; border-radius: 12px;
            font-size: 12px; font-weight: 600; }}
  .badge.PASS {{ background: #d1fae5; color: #065f46; }}
  .badge.FAIL, .badge.ERROR {{ background: #fee2e2; color: #991b1b; }}
  .badge.TIMEOUT {{ background: #fef3c7; color: #92400e; }}
  .query {{ font-weight: 600; }}
  .resp {{ white-space: pre-wrap; max-width: 520px; }}
  details {{ margin-top: 6px; }}
  details summary {{ cursor: pointer; color: #2563eb; font-size: 12px; }}
  pre {{ background: #f3f4f6; padding: 8px; border-radius: 6px; overflow-x: auto;
          font-size: 11px; max-width: 600px; }}
  .miss {{ color: #dc2626; }} .hit {{ color: #059669; }}
  .steps {{ margin-top: 8px; }}
  .steps summary {{ cursor: pointer; color: #2563eb; font-size: 12px; }}
  .step-line {{ font-size: 11px; padding: 2px 0; border-bottom: 1px dotted #eee;
                font-family: ui-monospace, SFMono-Regular, Menlo, monospace; }}
  .step-off {{ color: #6b7280; display: inline-block; min-width: 64px; }}
  .step-cat {{ display: inline-block; min-width: 110px; font-weight: 600; }}
  .cat-mcp_action {{ color: #7c3aed; }}
  .cat-knowledge_query {{ color: #0891b2; }}
  .cat-sql_query, .cat-dataverse_query {{ color: #b45309; }}
  .cat-planning {{ color: #4f46e5; }}
  .cat-oauth_card {{ color: #be185d; }}
  .cat-heartbeat {{ color: #9ca3af; }}
  .step-detail {{ color: #6b7280; font-size: 10px; white-space: pre-wrap; }}
</style>
</head>
<body>
<h1>Copilot Agent Eval Report</h1>
<div class="meta">Generated {generated}</div>
<div class="summary">
  <div class="card"><div class="n">{total}</div><div class="l">Total</div></div>
  <div class="card PASS"><div class="n">{passed}</div><div class="l">Passed</div></div>
  <div class="card FAIL"><div class="n">{failed}</div><div class="l">Failed</div></div>
  <div class="card TIMEOUT"><div class="n">{timeout}</div><div class="l">Timeout</div></div>
  <div class="card ERROR"><div class="n">{error}</div><div class="l">Error</div></div>
  <div class="card"><div class="n">{pass_rate}</div><div class="l">Pass rate</div></div>
</div>
<table>
<thead><tr>
  <th>Case</th><th>Status</th><th>Query</th><th>Response</th>
  <th>Latency</th><th>Assertions</th><th>Raw</th>
</tr></thead>
<tbody>
{rows}
</tbody>
</table>
</body></html>
"""


def _write_html(
    results: list[TestResult], path: Path, raw_activities: dict[str, list[dict]]
) -> None:
    total = len(results)
    passed = sum(1 for r in results if r.status == "PASS")
    failed = sum(1 for r in results if r.status == "FAIL")
    timeout = sum(1 for r in results if r.status == "TIMEOUT")
    error = sum(1 for r in results if r.status == "ERROR")
    pass_rate = f"{passed / total * 100:.1f}%" if total else "n/a"

    rows_html: list[str] = []
    for r in results:
        assertion_bits = []
        if r.keyword_hits:
            assertion_bits.append(
                f'<span class="hit">hit: {html.escape(", ".join(r.keyword_hits))}</span>'
            )
        if r.keyword_misses:
            assertion_bits.append(
                f'<span class="miss">missing: {html.escape(", ".join(r.keyword_misses))}</span>'
            )
        if r.error_phrases_hit:
            assertion_bits.append(
                f'<span class="miss">error phrases: {html.escape(", ".join(r.error_phrases_hit))}</span>'
            )
        if r.error:
            assertion_bits.append(f'<span class="miss">{html.escape(r.error)}</span>')
        assertion_html = "<br>".join(assertion_bits) or "—"

        raw_html = ""
        acts = raw_activities.get(r.case_id)
        if acts:
            raw_html = (
                f'<details><summary>view {len(acts)} activities</summary>'
                f'<pre>{html.escape(json.dumps(acts, indent=2, ensure_ascii=False))}</pre></details>'
            )

        # Steps timeline (MCP calls, knowledge queries, SQL, traces, etc.)
        steps_html = ""
        if r.steps:
            step_lines = []
            for s in r.steps:
                cat = s.category or ""
                name = f" [{html.escape(s.activity_name)}]" if s.activity_name else ""
                detail = ""
                if s.detail:
                    detail = f'<div class="step-detail">{html.escape(s.detail[:400])}</div>'
                step_lines.append(
                    f'<div class="step-line">'
                    f'<span class="step-off">+{s.offset_ms}ms</span> '
                    f'<span class="step-cat cat-{cat}">{html.escape(cat or s.activity_type)}</span> '
                    f'{html.escape(s.summary)}{name}{detail}</div>'
                )
            steps_html = (
                f'<details class="steps"><summary>{len(r.steps)} steps</summary>'
                f'{"".join(step_lines)}</details>'
            )

        latency = (
            f"{r.total_latency_ms} ms"
            if r.total_latency_ms is not None
            else "—"
        )
        rows_html.append(
            f"<tr>"
            f"<td><b>{html.escape(r.case_id)}</b></td>"
            f'<td><span class="badge {r.status}">{r.status}</span></td>'
            f'<td class="query">{html.escape(r.query)}</td>'
            f'<td class="resp">{html.escape(r.response_text or "")}{steps_html}</td>'
            f"<td>{html.escape(latency)}</td>"
            f"<td>{assertion_html}</td>"
            f"<td>{raw_html}</td>"
            f"</tr>"
        )

    page = _HTML_TEMPLATE.format(
        generated=datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        total=total,
        passed=passed,
        failed=failed,
        timeout=timeout,
        error=error,
        pass_rate=pass_rate,
        rows="\n".join(rows_html),
    )
    path.write_text(page, encoding="utf-8")
    log.info("HTML report written: %s", path)


# --------------------------------------------------------------------- JUnit
def _write_junit(results: list[TestResult], path: Path) -> None:
    suite = ET.Element(
        "testsuite",
        {
            "name": "Copilot Agent Eval",
            "tests": str(len(results)),
            "failures": str(sum(1 for r in results if r.status == "FAIL")),
            "errors": str(sum(1 for r in results if r.status in ("ERROR", "TIMEOUT"))),
            "timestamp": datetime.now().isoformat(),
        },
    )
    for r in results:
        case = ET.SubElement(
            suite,
            "testcase",
            {"classname": "copilot_agent_eval", "name": r.case_id, "time": f"{(r.total_latency_ms or 0) / 1000:.3f}"},
        )
        if r.status in ("FAIL", "ERROR", "TIMEOUT"):
            failure = ET.SubElement(
                case,
                "failure",
                {"message": r.error or f"Status: {r.status}", "type": r.status},
            )
            failure.text = (
                f"Query: {r.query}\nResponse: {r.response_text}\n"
                f"Missing keywords: {', '.join(r.keyword_misses)}\n"
                f"Error phrases: {', '.join(r.error_phrases_hit)}"
            )
    tree = ET.ElementTree(suite)
    ET.indent(tree, space="  ")
    tree.write(path, encoding="utf-8", xml_declaration=True)
    log.info("JUnit report written: %s", path)
